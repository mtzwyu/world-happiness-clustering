"""
Công cụ biên dịch Báo cáo Word — Đề tài 16 — Nhóm 12
Học phần: Khai thác dữ liệu (Data Mining)
Phụ trách: Hà Mạnh Trường (MSSV: 2045240326) — Mốc M4 (Ngày 23 – Ngày 25)

Chức năng:
1. Đọc bản gốc Markdown `reports/final_report/bao_cao.md`.
2. Dựng trang bìa học thuật (có viền khung) rồi ngắt trang sang phần nội dung.
3. Dựng mục lục tự động bằng trường TOC của Word.
4. Tự động đọc bảng từ `reports/tables/*.csv` theo cú pháp nhúng trong Markdown.
5. Tự động chèn hình 300 DPI từ `reports/figures/*.png` kèm chú thích.
6. Xuất đồng thời hai tệp Word: bao_cao.docx và Nhom12_BaoCao.docx.

Cú pháp nhúng trong Markdown:
    :::table
    file: model_comparison_metrics_by_year.csv
    caption: Bảng 4.1: Chỉ số đánh giá theo năm
    where: algorithm == 'K-Means' and scaling == 'standard'
    columns: year, silhouette, davies_bouldin
    round: 4
    :::

    :::figure
    file: radar_chart_clusters.png
    caption: Hình 4.3: Chân dung radar ba cụm
    width: 6.0
    :::

    :::toc
    :::

    :::pagebreak
    :::

Cách chạy:
    python reports/final_report/build_report.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# ĐƯỜNG DẪN TƯƠNG ĐỐI ĐỘC LẬP MÔI TRƯỜNG
# ---------------------------------------------------------------------------
PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "reports" / "final_report"
TABLES_DIR = PROJECT_ROOT / "reports" / "tables"
FIGURES_DIR = PROJECT_ROOT / "reports" / "figures"

SOURCE_MD = REPORT_DIR / "bao_cao.md"
OUTPUT_PRIMARY = REPORT_DIR / "bao_cao.docx"
OUTPUT_SUBMISSION = REPORT_DIR / "Nhom12_BaoCao.docx"

FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(13)
TABLE_SIZE = Pt(11)
CAPTION_SIZE = Pt(11.5)

HEADER_TEXT = "ĐỀ TÀI 16: PHÂN NHÓM QUỐC GIA THEO CÁC YẾU TỐ HẠNH PHÚC — NHÓM 12"

# Nội dung trang bìa (danh sách các dòng, kèm cỡ chữ và kiểu chữ)
COVER_BLOCK = [
    ("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH", 15, True, False),
    ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True, False),
    ("BỘ MÔN KHOA HỌC DỮ LIỆU", 14, True, False),
    ("", 12, False, False),
    ("BÁO CÁO ĐỒ ÁN MÔN HỌC", 13, False, False),
    ("KHAI THÁC DỮ LIỆU (DATA MINING)", 14, True, False),
    ("", 12, False, False),
    ("ĐỀ TÀI 16", 16, True, False),
    ("PHÂN NHÓM QUỐC GIA THEO CÁC YẾU TỐ HẠNH PHÚC", 16, True, False),
    ("(WORLD HAPPINESS REPORT 2015 – 2019)", 13, False, True),
    ("", 12, False, False),
    ("Giảng viên hướng dẫn: ThS. Đinh Nguyễn Trọng Nghĩa", 13, False, False),
    ("", 12, False, False),
    ("Sinh viên thực hiện:", 13, True, False),
    ("1. Hà Mạnh Trường — MSSV: 2045240326 — Nhóm trưởng (34%)", 13, False, False),
    ("2. Huỳnh Tâm Trí — MSSV: 2045240315 — Thành viên (33%)", 13, False, False),
    ("3. Ngô Quang Trung — MSSV: 2045240325 — Thành viên (33%)", 13, False, False),
    ("", 12, False, False),
    ("TP. Hồ Chí Minh, Năm 2026", 13, True, False),
]


# ===========================================================================
# 1. TIỆN ÍCH ĐỊNH DẠNG OOXML
# ===========================================================================
def set_run_font(run, size=None, bold=None, italic=None, color=None, name=FONT_NAME):
    """Gán font, cỡ chữ, kiểu chữ và màu cho một run, gồm cả thuộc tính eastAsia."""
    run.font.name = name
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attribute), name)


def force_style_font(style, name=FONT_NAME, size=None, bold=None, italic=None,
                     color=None, before=None, after=None, keep_with_next=False):
    """Áp font và khoảng cách cho một style, đồng thời gỡ thuộc tính font theo theme."""
    style.font.name = name
    if size is not None:
        style.font.size = size
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(theme_attr)) is not None:
            rfonts.attrib.pop(qn(theme_attr))
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attribute), name)

    if before is not None:
        style.paragraph_format.space_before = before
    if after is not None:
        style.paragraph_format.space_after = after
    if keep_with_next:
        style.paragraph_format.keep_with_next = True


def add_page_number_field(paragraph):
    """Chèn trường PAGE để Word tự đánh số trang."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def add_toc_field(paragraph):
    """Chèn trường TOC để Word tự dựng mục lục (cập nhật bằng F9)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Nhấn F9 (hoặc chuột phải chọn Update Field) để cập nhật mục lục."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def set_page_border(section, color="1B365D", size="24", space="24"):
    """Vẽ viền khung quanh trang cho một section (dùng cho trang bìa)."""
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), space)
        element.set(qn("w:color"), color)
        pg_borders.append(element)

    pg_mar = sect_pr.find(qn("w:pgMar"))
    if pg_mar is not None:
        pg_mar.addnext(pg_borders)
    else:
        sect_pr.append(pg_borders)


def set_cell_shading(cell, fill):
    """Tô nền cho một ô bảng, chèn đúng thứ tự phần tử theo schema OOXML."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    anchor = tc_pr.find(qn("w:vAlign"))
    if anchor is not None:
        anchor.addprevious(shading)
    else:
        tc_pr.append(shading)


def set_table_borders(table, color="7F8C8D"):
    """Vẽ đường viền mảnh cho bảng, chèn đúng thứ tự phần tử theo schema OOXML."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)

    anchor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        element = tbl_pr.find(qn(tag))
        if element is not None:
            anchor = element
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tbl_pr.append(borders)


def set_repeat_header(row):
    """Đánh dấu dòng tiêu đề bảng lặp lại trên mỗi trang in."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


# ===========================================================================
# 2. KHỞI TẠO DOCUMENT VÀ TRANG BÌA
# ===========================================================================
def configure_document() -> Document:
    """Tạo document và áp chuẩn định dạng in ấn (font, heading, lề trang)."""
    document = Document()

    force_style_font(document.styles["Normal"], size=BODY_SIZE)
    normal_pf = document.styles["Normal"].paragraph_format
    normal_pf.line_spacing = 1.3
    normal_pf.space_after = Pt(4)
    normal_pf.space_before = Pt(0)
    normal_pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    force_style_font(document.styles["Heading 1"], size=Pt(15), bold=True, italic=False,
                     color="1B365D", before=Pt(12), after=Pt(6), keep_with_next=True)
    force_style_font(document.styles["Heading 2"], size=Pt(13.5), bold=True, italic=False,
                     color="000000", before=Pt(8), after=Pt(4), keep_with_next=True)
    force_style_font(document.styles["Heading 3"], size=Pt(13), bold=True, italic=True,
                     color="000000", before=Pt(6), after=Pt(2), keep_with_next=True)
    force_style_font(document.styles["List Bullet"], size=BODY_SIZE)
    force_style_font(document.styles["List Number"], size=BODY_SIZE)
    document.styles["List Bullet"].paragraph_format.space_after = Pt(2)
    document.styles["List Number"].paragraph_format.space_after = Pt(2)

    for section in document.sections:
        apply_margins(section)
    return document


def apply_margins(section):
    """Canh lề chuẩn đóng gáy đồ án: trái 3,0 cm, phải 2,0 cm, trên/dưới 2,0 cm."""
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def build_cover(document: Document):
    """Dựng trang bìa học thuật có viền khung và ngắt trang sang phần nội dung."""
    for text, size, bold, italic in COVER_BLOCK:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.2
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=Pt(size), bold=bold, italic=italic, color="000000")

    # Viền khung cho section trang bìa
    set_page_border(document.sections[0])

    # Ngắt sang section mới cho phần nội dung (section break qua trang)
    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    apply_margins(body_section)

    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False

    header_paragraph = body_section.header.paragraphs[0]
    header_paragraph.text = ""
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_paragraph.add_run(HEADER_TEXT)
    set_run_font(header_run, size=Pt(10), italic=True, color="555555")

    footer_paragraph = body_section.footer.paragraphs[0]
    footer_paragraph.text = ""
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(footer_paragraph)
    for run in footer_paragraph.runs:
        set_run_font(run, size=Pt(11), color="000000")


# ===========================================================================
# 3. ĐỌC BẢNG TỪ CSV VÀ DỰNG BẢNG OOXML
# ===========================================================================
def load_csv_table(params: dict) -> tuple:
    """Đọc bảng từ reports/tables/*.csv theo tham số directive, trả về (header, rows, numeric)."""
    file_name = params.get("file")
    if not file_name:
        raise ValueError("Directive :::table thiếu tham số 'file'")
    csv_path = TABLES_DIR / file_name
    if not csv_path.exists():
        raise FileNotFoundError(f"Không tìm thấy bảng CSV: {csv_path}")

    frame = pd.read_csv(csv_path, encoding="utf-8")

    where = params.get("where")
    if where:
        frame = frame.query(where)

    columns = params.get("columns")
    if columns:
        selected = [column.strip() for column in columns.split(",") if column.strip()]
        frame = frame[selected]

    round_digits = params.get("round")
    round_digits = int(round_digits) if round_digits else None

    header = [str(column) for column in frame.columns]
    records = frame.to_dict("records")

    numeric_flags = []
    integer_flags = []
    for column in frame.columns:
        series = frame[column]
        numeric = pd.api.types.is_numeric_dtype(series)
        numeric_flags.append(numeric)
        non_null = series.dropna()
        integer_flags.append(
            bool(numeric)
            and len(non_null) > 0
            and bool((non_null.astype(float) % 1 == 0).all())
        )

    rows = []
    for record in records:
        row = []
        for index, column in enumerate(frame.columns):
            value = record[column]
            if pd.isna(value):
                row.append("")
            elif numeric_flags[index] and integer_flags[index]:
                row.append(str(int(round(float(value)))))
            elif numeric_flags[index]:
                digits = round_digits if round_digits is not None else 4
                row.append(f"{float(value):.{digits}f}".replace(".", ","))
            else:
                row.append(str(value))
        rows.append(row)

    return header, rows, numeric_flags


def parse_markdown_table(lines: list) -> tuple:
    """Phân tích bảng Markdown dạng ống (pipe table) thành (header, rows)."""
    def split_row(line):
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        return cells

    header = split_row(lines[0])
    body_lines = [line for line in lines[1:] if not re.match(r"^[\s|:-]+$", line)]
    rows = [split_row(line) for line in body_lines]
    numeric_flags = []
    for column_index in range(len(header)):
        values = [row[column_index] for row in rows if column_index < len(row)]
        numeric = bool(values) and all(
            re.match(r"^-?\d+([.,]\d+)?$", value.replace(" ", "")) for value in values
        )
        numeric_flags.append(numeric)
    return header, rows, numeric_flags


def render_table(document, caption, header, rows, numeric_flags):
    """Dựng bảng OOXML: tiêu đề trên bảng, header tô nền, cột số căn phải, viền mảnh."""
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_before = Pt(6)
        caption_paragraph.paragraph_format.space_after = Pt(2)
        caption_run = caption_paragraph.add_run(caption)
        set_run_font(caption_run, size=CAPTION_SIZE, bold=True, color="000000")

    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)

    header_cells = table.rows[0].cells
    for index, title in enumerate(header):
        header_cells[index].text = ""
        paragraph = header_cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(title))
        set_run_font(run, size=TABLE_SIZE, bold=True, color="000000")
        set_cell_shading(header_cells[index], "EAECEE")
    set_repeat_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if index < len(numeric_flags) and numeric_flags[index]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            set_run_font(run, size=TABLE_SIZE, color="000000")

    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def render_figure(document, params):
    """Chèn hình 300 DPI căn giữa kèm chú thích in nghiêng bên dưới."""
    file_name = params.get("file")
    if not file_name:
        raise ValueError("Directive :::figure thiếu tham số 'file'")
    image_path = FIGURES_DIR / file_name
    if not image_path.exists():
        raise FileNotFoundError(f"Không tìm thấy hình: {image_path}")

    width = float(params.get("width", 6.0))
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(4)
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width))

    caption = params.get("caption", "")
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_after = Pt(8)
        caption_run = caption_paragraph.add_run(caption)
        set_run_font(caption_run, size=CAPTION_SIZE, italic=True, color="000000")


# ===========================================================================
# 4. PHÂN TÍCH MARKDOWN VÀ DỰNG NỘI DUNG
# ===========================================================================
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline_runs(paragraph, text):
    """Thêm run cho đoạn văn, xử lý **đậm**, *nghiêng* và `mã`."""
    tokens = INLINE_PATTERN.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, italic=True)
        else:
            run = paragraph.add_run(token)
            set_run_font(run)


def parse_markdown(text: str) -> list:
    """Chuyển Markdown thành danh sách các block để renderer xử lý."""
    blocks = []
    lines = text.split("\n")
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith(":::"):
            directive = stripped[3:].strip()
            if directive in ("pagebreak", "toc"):
                index += 1
                while index < len(lines) and lines[index].strip() != ":::":
                    index += 1
                index += 1
                blocks.append({"type": directive})
                continue

            params = {"type": directive}
            index += 1
            while index < len(lines) and lines[index].strip() != ":::":
                if ":" in lines[index]:
                    key, _, value = lines[index].partition(":")
                    params[key.strip()] = value.strip()
                index += 1
            index += 1
            blocks.append(params)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append({"type": "heading", "level": min(level, 3), "text": stripped[level:].strip()})
            index += 1
            continue

        if re.match(r"^\|.*\|$", stripped):
            table_lines = []
            while index < len(lines) and re.match(r"^\|.*\|$", lines[index].strip()):
                table_lines.append(lines[index].strip())
                index += 1
            blocks.append({"type": "md_table", "lines": table_lines})
            continue

        if re.match(r"^[-*]\s+", stripped):
            items = []
            while index < len(lines) and re.match(r"^[-*]\s+", lines[index].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[index].strip()))
                index += 1
            blocks.append({"type": "bullets", "items": items})
            continue

        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[index].strip()))
                index += 1
            blocks.append({"type": "numbers", "items": items})
            continue

        # Đoạn văn thường: gom các dòng liền nhau đến khi gặp block khác
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", ":::", "|"))
                or re.match(r"^[-*]\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines)})

    return blocks


def render_blocks(document: Document, blocks: list):
    """Render danh sách block ra document Word."""
    heading_count = {"h1": 0, "h2": 0, "h3": 0}
    for block in blocks:
        block_type = block["type"]

        if block_type == "heading":
            level = block["level"]
            key = f"h{level}"
            heading_count[key] += 1
            paragraph = document.add_heading(level=level)
            add_inline_runs(paragraph, block["text"])
            for run in paragraph.runs:
                if level == 1:
                    set_run_font(run, size=Pt(15), bold=True, color="1B365D")
                elif level == 2:
                    set_run_font(run, size=Pt(13.5), bold=True, color="000000")
                else:
                    set_run_font(run, size=Pt(13), bold=True, italic=True, color="000000")

        elif block_type == "paragraph":
            paragraph = document.add_paragraph()
            add_inline_runs(paragraph, block["text"])

        elif block_type == "bullets":
            for item in block["items"]:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline_runs(paragraph, item)

        elif block_type == "numbers":
            for item in block["items"]:
                paragraph = document.add_paragraph(style="List Number")
                add_inline_runs(paragraph, item)

        elif block_type == "md_table":
            header, rows, numeric_flags = parse_markdown_table(block["lines"])
            render_table(document, None, header, rows, numeric_flags)

        elif block_type == "table":
            header, rows, numeric_flags = load_csv_table(block)
            render_table(document, block.get("caption"), header, rows, numeric_flags)

        elif block_type == "figure":
            render_figure(document, block)

        elif block_type == "toc":
            toc_paragraph = document.add_paragraph()
            add_toc_field(toc_paragraph)

        elif block_type == "pagebreak":
            document.add_page_break()


# ===========================================================================
# 5. HÀM CHÍNH
# ===========================================================================
def build_report():
    """Biên dịch bao_cao.md thành hai tệp Word đầu ra."""
    if not SOURCE_MD.exists():
        raise FileNotFoundError(f"Không tìm thấy bản gốc Markdown: {SOURCE_MD}")

    markdown_text = SOURCE_MD.read_text(encoding="utf-8")
    blocks = parse_markdown(markdown_text)

    document = configure_document()
    build_cover(document)
    render_blocks(document, blocks)

    OUTPUT_PRIMARY.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT_PRIMARY))
    document.save(str(OUTPUT_SUBMISSION))

    print("=== BIÊN DỊCH BÁO CÁO WORD ===")
    print(f"Bản gốc Markdown     : {SOURCE_MD}")
    print(f"Số block nội dung    : {len(blocks)}")
    print(f"Đã xuất tệp          : {OUTPUT_PRIMARY}")
    print(f"Đã xuất tệp          : {OUTPUT_SUBMISSION}")
    print(f"Dung lượng bao_cao.docx      : {OUTPUT_PRIMARY.stat().st_size} bytes")
    print(f"Dung lượng Nhom12_BaoCao.docx: {OUTPUT_SUBMISSION.stat().st_size} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(build_report())
